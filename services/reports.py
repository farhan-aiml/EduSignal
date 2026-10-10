"""Reports built only from stored database data, exported as formatted Excel or PDF.

Scopes and permissions (checked here, not in the page):
  student -> only their own record
  teacher -> one of their students, or one of their classes
  admin   -> any student, any class, or the whole institution
Every export is written to the audit log. Empty data produces an explicit "no data" report, never invented rows.
"""
from __future__ import annotations

import io
from datetime import datetime

import pandas as pd

from database import database as db
from services import audit
from services.access import PermissionDenied, can_manage_section, can_view_student, teacher_sections

REPORTS = {
    "attendance": "Attendance Report",
    "marks": "Subject-wise Marks Report",
    "performance": "Student Performance Report",
    "prediction": "AI Prediction Report",
    "study_plan": "Study Plan Progress Report",
}


# ---------------- scope ----------------
def allowed_scopes(user: dict) -> list[str]:
    return {"student": ["student"], "teacher": ["student", "class"], "admin": ["student", "class", "institution"]}.get(
        user.get("role"), [])


def resolve_students(user: dict, scope: str, value: str | None) -> list[str]:
    """Student ids covered by a report, after the permission check."""
    if scope not in allowed_scopes(user):
        raise PermissionDenied("Aapko yeh report dekhne ki permission nahi hai.")
    if scope == "student":
        if user.get("role") == "student":
            value = user.get("student_id")
        if not value or not can_view_student(user, value):
            raise PermissionDenied("Aapko is student ki report download karne ki permission nahi hai.")
        return [value]
    if scope == "class":
        if not value or not can_manage_section(user, value):
            raise PermissionDenied("Aapko is class ki report download karne ki permission nahi hai.")
        return db.query_df("SELECT student_id FROM students WHERE class_section = ? ORDER BY student_id", (value,))[
            "student_id"].tolist()
    return db.query_df("SELECT student_id FROM students ORDER BY student_id")["student_id"].tolist()


def scope_options(user: dict, scope: str) -> list[str]:
    if scope == "class":
        secs = db.class_sections() if user.get("role") == "admin" else teacher_sections(user["username"])
        return list(secs)
    if scope == "student":
        if user.get("role") == "student":
            return [user.get("student_id")]
        if user.get("role") == "admin":
            return db.students_df()["student_id"].tolist()
        secs = teacher_sections(user["username"])
        return db.students_df().query("class_section in @secs")["student_id"].tolist()
    return []


# ---------------- data ----------------
def _in(ids):
    return ",".join("?" * len(ids)) or "''"


def attendance_data(ids):
    return db.query_df(
        "SELECT st.student_id AS 'Student ID', st.name AS 'Name', st.class_section AS 'Class', "
        "COALESCE(sub.name, s.subject_code) AS 'Subject', SUM(r.status IN ('present','late')) AS 'Attended', "
        "SUM(r.status != 'excused') AS 'Held', SUM(r.status = 'absent') AS 'Absent', SUM(r.status = 'late') AS 'Late', "
        "SUM(r.status = 'excused') AS 'Excused', "
        "ROUND(100.0 * SUM(r.status IN ('present','late')) / NULLIF(SUM(r.status != 'excused'), 0), 1) AS 'Attendance %' "
        "FROM attendance_records r JOIN attendance_sessions s ON s.id = r.session_id "
        "JOIN students st ON st.student_id = r.student_id LEFT JOIN subjects sub ON sub.code = s.subject_code "
        f"WHERE r.student_id IN ({_in(ids)}) GROUP BY 1, 2, 3, 4 ORDER BY 1, 4", tuple(ids))


def marks_data(ids):
    return db.query_df(
        "SELECT st.student_id AS 'Student ID', st.name AS 'Name', COALESCE(sub.name, a.subject_code) AS 'Subject', "
        "a.title AS 'Assessment', a.kind AS 'Type', a.assessed_on AS 'Date', m.marks AS 'Marks', a.max_marks AS 'Out of', "
        "CASE WHEN m.absent = 1 THEN 'Absent' WHEN m.marks IS NULL THEN 'Missing' "
        "ELSE CAST(ROUND(100.0 * m.marks / a.max_marks, 1) AS TEXT) END AS '%' "
        "FROM assessments a JOIN students st ON st.class_section = a.class_section "
        "LEFT JOIN assessment_marks m ON m.assessment_id = a.id AND m.student_id = st.student_id "
        "LEFT JOIN subjects sub ON sub.code = a.subject_code "
        f"WHERE st.student_id IN ({_in(ids)}) ORDER BY 1, 3, a.assessed_on", tuple(ids))


def performance_data(ids):
    rec = db.records_df()
    rec = rec[rec["student_id"].isin(ids)]
    if rec.empty:
        return pd.DataFrame()
    rows = []
    for sid, g in rec.groupby("student_id"):
        g = g.sort_values(["record_date", "id"])
        last, first = g.iloc[-1], g.iloc[0]
        rows.append({"Student ID": sid, "Name": last["name"], "Class": last["class_section"], "Records": len(g),
                     "Latest record": last["record_date"], "Attendance %": last["Attendance"],
                     "Assignment": last["Assignment_Score"], "Internal": last["Internal_Marks"],
                     "Study h/day": last["Study_Hours"], "Previous marks": last["Previous_Marks"],
                     "Actual final score": last["final_score"] if pd.notna(last["final_score"]) else None,
                     "Change in attendance": round(float(last["Attendance"] - first["Attendance"]), 1)})
    return pd.DataFrame(rows)


def prediction_data(ids):
    return db.query_df(
        "SELECT p.student_id AS 'Student ID', st.name AS 'Name', p.created_at AS 'Predicted on', "
        "ROUND(p.predicted_score, 1) AS 'Predicted score', p.grade AS 'Grade', "
        "ROUND(p.interval_low, 1) AS '80% low', ROUND(p.interval_high, 1) AS '80% high', "
        "ra.risk_level AS 'Risk', ROUND(ra.risk_score, 0) AS 'Risk score', p.model_name AS 'Model', "
        "p.model_version AS 'Model version' FROM predictions p JOIN students st ON st.student_id = p.student_id "
        "LEFT JOIN risk_assessments ra ON ra.prediction_id = p.id "
        f"WHERE p.student_id IN ({_in(ids)}) ORDER BY p.student_id, p.id DESC", tuple(ids))


def study_plan_data(ids):
    return db.query_df(
        "SELECT t.student_id AS 'Student ID', st.name AS 'Name', t.plan_date AS 'Date', COUNT(*) AS 'Tasks', "
        "SUM(t.status = 'completed') AS 'Completed', SUM(t.status = 'in_progress') AS 'In progress', "
        "SUM(t.status = 'pending') AS 'Pending', ROUND(100.0 * SUM(t.status = 'completed') / COUNT(*), 0) AS 'Progress %', "
        "SUM(t.minutes) AS 'Planned min', SUM(CASE WHEN t.status = 'completed' THEN t.minutes ELSE 0 END) AS 'Done min' "
        "FROM study_tasks t JOIN students st ON st.student_id = t.student_id "
        f"WHERE t.student_id IN ({_in(ids)}) GROUP BY 1, 2, 3 ORDER BY 1, 3 DESC", tuple(ids))


BUILDERS = {"attendance": attendance_data, "marks": marks_data, "performance": performance_data,
            "prediction": prediction_data, "study_plan": study_plan_data}


def build(user: dict, report: str, scope: str, value: str | None) -> tuple[pd.DataFrame, dict]:
    if report not in BUILDERS:
        raise ValueError("Yeh report available nahi hai.")
    ids = resolve_students(user, scope, value)
    if scope == "student":
        value = ids[0]          # a student always gets their own id, whatever was requested
    df = BUILDERS[report](ids) if ids else pd.DataFrame()
    meta = {"title": REPORTS[report], "scope": {"student": f"Student {value}", "class": f"Class {value}",
                                                "institution": "Poora institution"}[scope],
            "generated_by": user.get("display_name") or user.get("username"),
            "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M"), "rows": len(df), "students": len(ids)}
    return df, meta


# ---------------- export ----------------
def to_excel(df: pd.DataFrame, meta: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = meta["title"][:31]
    ws["A1"] = f"EduSignal - {meta['title']}"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = f"{meta['scope']} | Generated by {meta['generated_by']} on {meta['generated_at']} | {meta['rows']} rows"
    ws["A2"].font = Font(italic=True, color="555555")
    if df.empty:
        ws["A4"] = "Is report ke liye database mein abhi koi data nahi hai."
    else:
        head = 4
        for j, col in enumerate(df.columns, 1):
            c = ws.cell(row=head, column=j, value=str(col))
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F2A30")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, row in enumerate(df.itertuples(index=False), head + 1):
            for j, v in enumerate(row, 1):
                ws.cell(row=i, column=j, value=None if pd.isna(v) else (v.item() if hasattr(v, "item") else v))
        for j, col in enumerate(df.columns, 1):
            width = max([len(str(col))] + [len(str(v)) for v in df.iloc[:200, j - 1].tolist()])
            ws.column_dimensions[get_column_letter(j)].width = min(45, max(10, width + 2))
        ws.freeze_panes = ws.cell(row=head + 1, column=1)
        ws.auto_filter.ref = f"A{head}:{get_column_letter(len(df.columns))}{head + len(df)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(df: pd.DataFrame, meta: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    buf = io.BytesIO()
    wide = len(df.columns) > 6
    page = landscape(A4) if wide else A4
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=14 * mm,
                            bottomMargin=14 * mm, title=f"EduSignal - {meta['title']}")
    ss = getSampleStyleSheet()
    cell = ss["BodyText"].clone("cell", fontSize=7.5 if wide else 8.5, leading=9.5 if wide else 10.5)
    story = [Paragraph(f"EduSignal - {meta['title']}", ss["Title"]),
             Paragraph(f"{meta['scope']} &nbsp;|&nbsp; Generated by {meta['generated_by']} on {meta['generated_at']}"
                       f" &nbsp;|&nbsp; {meta['rows']} rows", ss["Normal"]), Spacer(1, 6 * mm)]
    if df.empty:
        story.append(Paragraph("Is report ke liye database mein abhi koi data nahi hai.", ss["Normal"]))
    else:
        from xml.sax.saxutils import escape
        fmt = lambda v: "" if v is None or (isinstance(v, float) and pd.isna(v)) else escape(f"{v:g}" if isinstance(v, float) else str(v))
        head = cell.clone("hcell", textColor=colors.white)
        data = [[Paragraph(f"<b>{escape(str(c))}</b>", head) for c in df.columns]]
        data += [[Paragraph(fmt(v), cell) for v in row] for row in df.itertuples(index=False)]
        avail = page[0] - 24 * mm
        t = Table(data, repeatRows=1, colWidths=[avail / len(df.columns)] * len(df.columns))
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2A30")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F5F7")]),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#C9D1D9")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(t)
    story += [Spacer(1, 5 * mm), Paragraph("Yeh report EduSignal database ke stored records se bani hai. Predictions "
                                           "andaaze hain, final faisla teacher ka hai.", ss["Italic"])]
    doc.build(story)
    return buf.getvalue()


def export(user: dict, report: str, scope: str, value: str | None, fmt: str) -> tuple[bytes, str, str, dict]:
    """Returns (file bytes, file name, mime type, meta). Permission is enforced inside build()."""
    df, meta = build(user, report, scope, value)
    stem = f"edusignal_{report}_{(value or scope).replace(' ', '_')}_{datetime.now():%Y%m%d}"
    if fmt == "xlsx":
        data, name, mime = to_excel(df, meta), stem + ".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif fmt == "pdf":
        data, name, mime = to_pdf(df, meta), stem + ".pdf", "application/pdf"
    else:
        raise ValueError("Format sirf PDF ya Excel ho sakta hai.")
    audit.log("report.export", user, "report", report, {"scope": scope, "value": value, "format": fmt, "rows": meta["rows"]})
    return data, name, mime, meta
